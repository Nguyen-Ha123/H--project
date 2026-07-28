#!/usr/bin/env python3
"""Generate ADD-ON TCs for Quan_ly_cho_vay: Sửa/Xóa (sheet list) + Process bar (chi tiết).
New rows highlighted yellow. Does not modify existing TC file (unavailable in cloud).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

YELLOW = PatternFill("solid", fgColor="FFFF99")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6EAF8")
SECTION_FONT = Font(bold=True, color="1F4E79")
NOTE_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

COLS = [
    "ID_UC",
    "ID_TC",
    "Tiêu đề",
    "Tiền điều kiện",
    "Các bước thực hiện",
    "Kết quả mong muốn",
    "Kết quả test",
    "Ghi chú",
    "Flag",
]

UC = "Quan_ly_cho_vay"
NAV = (
    "Đăng nhập hệ thống → vào màn hình Danh sách hồ sơ Uỷ thác cho vay "
    "(sheet đầu / danh sách hồ sơ)"
)
NAV_DETAIL = (
    "Đăng nhập → Danh sách hồ sơ Uỷ thác cho vay → mở Chi tiết 1 hồ sơ "
    "đang ở trạng thái tương ứng"
)

PROCESS_STEPS = [
    "1. Chờ tiếp nhận",
    "2. Chờ phân công",
    "3. Chờ thẩm định",
    "4. Chờ ký duyệt",
    "5. Đã ký duyệt",
]


def rows():
    """Return list of dict rows. Section headers have ID_TC empty and Flag=SECTION."""
    out = []

    def section(title):
        out.append(
            {
                "ID_UC": UC,
                "ID_TC": "",
                "Tiêu đề": title,
                "Tiền điều kiện": "",
                "Các bước thực hiện": "",
                "Kết quả mong muốn": "",
                "Kết quả test": "",
                "Ghi chú": "SECTION — TC mới (bôi vàng)",
                "Flag": "SECTION",
            }
        )

    def tc(tid, title, pre, steps, expected, note=""):
        out.append(
            {
                "ID_UC": UC,
                "ID_TC": tid,
                "Tiêu đề": title,
                "Tiền điều kiện": pre,
                "Các bước thực hiện": steps,
                "Kết quả mong muốn": expected,
                "Kết quả test": "",  # empty — do not overwrite old results
                "Ghi chú": note or "NEW — merge vào file gốc, không đụng TC cũ",
                "Flag": "NEW",
            }
        )

    # ------------------------------------------------------------------
    # SHEET ĐẦU: SỬA HỒ SƠ
    # ------------------------------------------------------------------
    section("Sửa hồ sơ (Danh sách — sheet đầu)")

    tc(
        "TC_EDIT_001",
        "Hiển thị action Sửa trên danh sách hồ sơ",
        f"{NAV}. Có ≥1 hồ sơ trên lưới.",
        "1. Quan sát cột Thao tác / action trên từng dòng hồ sơ\n"
        "2. Hover / kiểm tra icon hoặc nút Sửa",
        "Mỗi dòng hồ sơ hợp lệ hiển thị action Sửa (icon/bút chì hoặc menu).\n"
        "Action nhận diện được, không bị che.",
    )

    tc(
        "TC_EDIT_002",
        "Mở form Sửa Hồ sơ - Uỷ thác cho vay từ danh sách",
        f"{NAV}. Có ≥1 hồ sơ.",
        "1. Tại dòng hồ sơ bất kỳ, click Sửa\n"
        "2. Quan sát popup/form mở ra",
        "Mở form tiêu đề: \"Sửa Hồ sơ - Uỷ thác cho vay\".\n"
        "Có nút đóng (X), Hủy, Lưu.\n"
        "Form load đủ section theo UI (Hình thức tiếp nhận, Người nộp, "
        "Thông tin tờ khai, Dự án, Đề nghị vay, Giải ngân, Thẩm định/Tài sản, "
        "Hồ sơ đính kèm).",
    )

    tc(
        "TC_EDIT_003",
        "Form Sửa prefill đúng dữ liệu hồ sơ đã lưu",
        f"{NAV}. Biết sẵn giá trị các field của 1 hồ sơ (gồm phần tài sản).",
        "1. Click Sửa trên hồ sơ đã chọn\n"
        "2. So sánh từng field trên form với dữ liệu đã lưu / chi tiết hồ sơ",
        "Toàn bộ field hiển thị đúng giá trị hiện tại của hồ sơ "
        "(không blank sai, không lệch sang hồ sơ khác).\n"
        "Gồm cả field tài sản: Tài sản đảm bảo, Mô tả chi tiết TSĐB, "
        "Hiện trạng tài sản, Tính pháp lý tài sản.",
    )

    tc(
        "TC_EDIT_004",
        "Sửa — Hủy không lưu thay đổi",
        f"{NAV}. Mở form Sửa.",
        "1. Đổi giá trị 1 field bất kỳ (vd: Họ tên người nộp)\n"
        "2. Click Hủy\n"
        "3. Mở lại Sửa / Chi tiết cùng hồ sơ",
        "Form đóng, không lưu.\n"
        "Dữ liệu hồ sơ giữ nguyên giá trị cũ.",
    )

    tc(
        "TC_EDIT_005",
        "Sửa — đóng icon X không lưu",
        f"{NAV}. Mở form Sửa.",
        "1. Đổi 1 field\n"
        "2. Click icon X\n"
        "3. Kiểm tra lại dữ liệu hồ sơ",
        "Form đóng, không lưu. Dữ liệu gốc không đổi.",
    )

    tc(
        "TC_EDIT_006",
        "Sửa — bỏ trống field bắt buộc (Người nộp / Tổ chức)",
        f"{NAV}. Mở form Sửa. Các field (*) đang có giá trị.",
        "1. Xóa trống lần lượt các field (*): "
        "Loại giấy tờ người nộp / Số định danh / Họ tên người nộp / "
        "Mã số thuế / các field (*) khác visible\n"
        "2. Click Lưu",
        "Không lưu thành công.\n"
        "Highlight field lỗi + message kiểu \"Vui lòng nhập + tên trường\" "
        "(theo UI thực tế).\n"
        "Nút Lưu disabled hoặc form vẫn mở. [TBD message nguyên văn nếu SRS khác]",
        "Assumption: rule required giống form Tạo mới",
    )

    tc(
        "TC_EDIT_007",
        "Sửa phần Tài sản đảm bảo — happy path",
        f"{NAV}. Hồ sơ cho phép sửa. Mở form Sửa → section "
        "\"5. Kết quả thẩm định hồ sơ của ngân hàng\".",
        "1. Sửa \"Tài sản đảm bảo\" = giá trị hợp lệ mới\n"
        "2. Sửa \"Mô tả chi tiết tài sản đảm bảo\"\n"
        "3. Sửa \"Hiện trạng tài sản\"\n"
        "4. Sửa \"Tính pháp lý tài sản\"\n"
        "5. Giữ checkbox cam kết (nếu bắt buộc)\n"
        "6. Click Lưu\n"
        "7. Mở Chi tiết / Sửa lại để verify",
        "Lưu thành công (message theo UI, vd: \"Cập nhật thành công\" — [TBD]).\n"
        "4 field tài sản phản ánh đúng giá trị mới trên Chi tiết hồ sơ.\n"
        "Các field khác không bị ghi đè sai.",
    )

    tc(
        "TC_EDIT_008",
        "Sửa phần Tài sản — bỏ trống field (*) tài sản",
        f"{NAV}. Mở form Sửa, section thẩm định/tài sản.",
        "1. Xóa trống lần lượt: Tài sản đảm bảo / Mô tả chi tiết TSĐB / "
        "Hiện trạng tài sản / Tính pháp lý tài sản\n"
        "2. Click Lưu sau mỗi lần (hoặc cuối cùng)",
        "Không cho lưu khi field (*) tài sản trống.\n"
        "Hiển thị validation + highlight field.\n"
        "Message: \"Vui lòng nhập + tên trường\" (ưu tiên text UI).",
    )

    tc(
        "TC_EDIT_009",
        "Sửa — cập nhật Hình thức tiếp nhận (Trực tiếp / Trực tuyến)",
        f"{NAV}. Mở form Sửa.",
        "1. Đổi radio Hình thức từ Trực tiếp ↔ Trực tuyến\n"
        "2. Click Lưu\n"
        "3. Mở lại form / chi tiết",
        "Lưu thành công. Hình thức hiển thị đúng lựa chọn mới.",
    )

    tc(
        "TC_EDIT_010",
        "Sửa — không cho sửa / ẩn Sửa theo trạng thái hồ sơ (nếu có rule)",
        f"{NAV}. Có hồ sơ ở các trạng thái: Chờ tiếp nhận, Chờ phân công, "
        "Chờ thẩm định, Chờ ký duyệt, Đã ký duyệt.",
        "1. Với từng trạng thái, kiểm tra action Sửa trên danh sách\n"
        "2. Nếu mở được form: thử Lưu thay đổi nhỏ",
        "Theo business rule [TBD từ SRS]:\n"
        "- Trạng thái cho phép sửa → action Sửa available, lưu OK\n"
        "- Trạng thái không cho sửa (thường: Đã ký duyệt / đã xử lý) → "
        "ẩn/disable Sửa hoặc Lưu bị chặn + message rõ\n"
        "Ghi nhận thực tế UI vào Kết quả test.",
        "BLOCKED/TBD — cần confirm rule theo trạng thái từ SRS",
    )

    # ------------------------------------------------------------------
    # SHEET ĐẦU: XÓA HỒ SƠ
    # ------------------------------------------------------------------
    section("Xóa hồ sơ (Danh sách — sheet đầu)")

    tc(
        "TC_DEL_001",
        "Hiển thị action Xóa trên danh sách hồ sơ",
        f"{NAV}. Có ≥1 hồ sơ.",
        "1. Quan sát cột Thao tác trên lưới\n"
        "2. Kiểm tra icon/nút Xóa",
        "Action Xóa hiển thị trên dòng hồ sơ (theo quyền + trạng thái cho phép).",
    )

    tc(
        "TC_DEL_002",
        "Xóa hồ sơ — hủy trên dialog xác nhận",
        f"{NAV}. Chọn 1 hồ sơ còn trên list.",
        "1. Click Xóa\n"
        "2. Quan sát dialog xác nhận\n"
        "3. Chọn Hủy / Không / Đóng",
        "Hiện dialog xác nhận xóa (nội dung rõ mã/tên hồ sơ nếu có).\n"
        "Sau khi Hủy: dialog đóng, hồ sơ vẫn còn trên danh sách, dữ liệu không đổi.",
    )

    tc(
        "TC_DEL_003",
        "Xóa hồ sơ — xác nhận thành công",
        f"{NAV}. Có hồ sơ test chuyên dụng (không ảnh hưởng data thật).",
        "1. Ghi nhận mã/ID hồ sơ trước khi xóa\n"
        "2. Click Xóa → Xác nhận\n"
        "3. Quan sát danh sách + (nếu có) tìm lại theo mã",
        "Message thành công theo UI (vd: \"Xóa thành công\" — [TBD]).\n"
        "Hồ sơ biến mất khỏi danh sách.\n"
        "Search lại mã → Không có dữ liệu / không tìm thấy.",
    )

    tc(
        "TC_DEL_004",
        "Xóa hồ sơ — ràng buộc theo trạng thái",
        f"{NAV}. Có hồ sơ từng trạng thái process "
        "(Chờ tiếp nhận … Đã ký duyệt).",
        "1. Thử Xóa với từng trạng thái\n"
        "2. Ghi nhận cho phép / chặn",
        "Theo rule [TBD]:\n"
        "- Trạng thái cho phép xóa → xóa OK như TC_DEL_003\n"
        "- Trạng thái không cho xóa (vd Đã ký duyệt) → ẩn/disable Xóa "
        "hoặc báo lỗi, hồ sơ vẫn còn\n"
        "Ưu tiên behavior UI thực tế.",
        "TBD rule theo trạng thái",
    )

    tc(
        "TC_DEL_005",
        "Xóa hồ sơ — kiểm tra phần tài sản/thẩm định bị gỡ theo hồ sơ",
        f"{NAV}. Hồ sơ có dữ liệu section tài sản đã lưu; trạng thái cho phép xóa.",
        "1. Mở Chi tiết xác nhận có dữ liệu Tài sản đảm bảo / TSĐB\n"
        "2. Quay lại danh sách → Xóa → Xác nhận\n"
        "3. Cố mở lại chi tiết / query theo ID (nếu có quyền/API/DB)",
        "Hồ sơ và dữ liệu liên quan (gồm phần tài sản/thẩm định) không còn "
        "truy cập được trên UI.\n"
        "Không còn orphan hiển thị trên list. [TBD SQL table nếu SRS cung cấp]",
    )

    # ------------------------------------------------------------------
    # CHI TIẾT: THANH PROCESS
    # ------------------------------------------------------------------
    section("Chi tiết hồ sơ — Thanh process (highlight theo trạng thái)")

    tc(
        "TC_PROC_001",
        "Thanh process hiển thị đủ 5 bước trạng thái",
        f"{NAV_DETAIL}.",
        "1. Mở tab CHI TIẾT HỒ SƠ\n"
        "2. Quan sát thanh process phía trên",
        "Thanh process hiển thị đúng 5 bước theo thứ tự:\n"
        + "\n".join(f"- {s}" for s in PROCESS_STEPS)
        + "\nCác bước nối nhau trên 1 thanh ngang, nhìn thấy số thứ tự.",
    )

    tc(
        "TC_PROC_002",
        "Highlight process = trạng thái hồ sơ hiện tại (mapping)",
        f"{NAV_DETAIL}. Biết trạng thái hồ sơ trên list / header.",
        "1. Đọc trạng thái hồ sơ (list hoặc nhãn trên chi tiết)\n"
        "2. So với bước đang được highlight trên thanh process",
        "Bước được highlight / active trên thanh process KHỚP đúng "
        "trạng thái hiện tại của hồ sơ.\n"
        "Không highlight sai bước.",
    )

    # One TC per status — continuation style for clarity as separate TCs
    status_cases = [
        ("TC_PROC_003", "Chờ tiếp nhận", 1),
        ("TC_PROC_004", "Chờ phân công", 2),
        ("TC_PROC_005", "Chờ thẩm định", 3),
        ("TC_PROC_006", "Chờ ký duyệt", 4),
        ("TC_PROC_007", "Đã ký duyệt", 5),
    ]
    for tid, status, step in status_cases:
        tc(
            tid,
            f"Process highlight khi trạng thái = {status}",
            f"{NAV_DETAIL.replace('trạng thái tương ứng', f'trạng thái \"{status}\"')}.",
            f"1. Mở Chi tiết hồ sơ trạng thái \"{status}\"\n"
            f"2. Quan sát bước số {step} trên thanh process\n"
            f"3. Quan sát các bước trước/sau",
            f"Bước {step} \"{status}\" được highlight (active / màu nhấn — "
            f"theo UI, vd chữ xanh đậm hoặc vòng active).\n"
            f"Các bước < {step}: thể hiện đã qua (completed style).\n"
            f"Các bước > {step}: chưa tới (inactive/xám) — trừ khi UI khác "
            f"(ghi nhận thực tế).\n"
            f"Trạng thái hồ sơ = \"{status}\".",
        )

    tc(
        "TC_PROC_008",
        "Đồng bộ trạng thái List ↔ Process bar Chi tiết",
        f"{NAV}. Có ≥1 hồ sơ mỗi trạng thái (hoặc sample đa trạng thái).",
        "1. Từ danh sách, ghi trạng thái cột Trạng thái của hồ sơ A\n"
        "2. Click vào hồ sơ A → Chi tiết\n"
        "3. Đối chiếu bước highlight trên process bar",
        "Trạng thái trên list = bước đang highlight trên process bar chi tiết.\n"
        "Không lệch (vd list \"Chờ thẩm định\" mà process highlight \"Đã ký duyệt\").",
    )

    tc(
        "TC_PROC_009",
        "Thanh process chỉ đọc — không đổi trạng thái bằng click bước",
        f"{NAV_DETAIL}.",
        "1. Click lần lượt các bước trên thanh process\n"
        "2. Quan sát trạng thái hồ sơ / highlight",
        "Click bước không làm thay đổi trạng thái hồ sơ.\n"
        "Highlight vẫn theo trạng thái hiện tại.\n"
        "Không có navigation/action ngoài ý muốn (trừ khi SRS định nghĩa — [TBD]).",
    )

    tc(
        "TC_PROC_010",
        "Process bar vẫn hiển thị khi chuyển tab chi tiết",
        f"{NAV_DETAIL}.",
        "1. Ở CHI TIẾT HỒ SƠ — xác nhận process bar\n"
        "2. Chuyển tab THÔNG TIN THẨM ĐỊNH HỒ SƠ\n"
        "3. Chuyển tab QUÁ TRÌNH XỬ LÝ\n"
        "4. Quay lại CHI TIẾT HỒ SƠ",
        "Thanh process vẫn đúng trạng thái (không reset / không mất highlight).\n"
        "Behavior theo UI: process bar sticky trên các tab hoặc luôn đúng khi "
        "quay lại chi tiết — ghi nhận thực tế.",
    )

    tc(
        "TC_PROC_011",
        "Sau khi cập nhật trạng thái (workflow) — process bar highlight bước mới",
        "Có quyền đẩy trạng thái hồ sơ (tiếp nhận / phân công / thẩm định / "
        "ký duyệt) theo luồng hệ thống. [TBD tên action thực tế]",
        "1. Mở chi tiết hồ sơ ở trạng thái N\n"
        "2. Thực hiện action chuyển sang trạng thái N+1 (theo UI)\n"
        "3. Quan sát thanh process",
        "Highlight chuyển sang bước tương ứng trạng thái mới.\n"
        "Bước cũ chuyển style completed.\n"
        "List cập nhật cùng trạng thái mới.",
        "Assumption: có action chuyển trạng thái trên UI",
    )

    return out


def style_cell(cell, fill=None, font=None):
    cell.alignment = WRAP
    cell.border = THIN
    if fill:
        cell.fill = fill
    if font:
        cell.font = font


def build(path: Path):
    wb = Workbook()

    # --- MERGE note sheet ---
    ws0 = wb.active
    ws0.title = "HUONG_DAN_MERGE"
    notes = [
        ["QUAN TRỌNG — Merge TC mới vào file gốc"],
        [""],
        ["File gốc (máy local): C:\\Users\\hant2\\Desktop\\HATC\\Output\\Quan_ly_cho_vay_ fix comment"],
        ["Cloud agent KHÔNG đọc được file gốc → xuất sheet TC_MOI (bôi vàng)."],
        [""],
        ["Cách merge:"],
        ["1. Mở file gốc Excel (giữ nguyên mọi TC + cột Kết quả test cũ)."],
        ["2. Copy toàn bộ dòng Flag=NEW từ sheet TC_MOI."],
        ["3. Paste vào sheet đầu (Danh sách / Sửa-Xóa) và sheet Chi tiết tương ứng."],
        ["4. KHÔNG ghi đè / xóa TC cũ; KHÔNG sửa cột Kết quả test đã có."],
        ["5. Đổi ID_TC cho khớp sequence file gốc nếu team đánh số liên tục."],
        ["6. Giữ tô vàng các dòng mới để review."],
        [""],
        ["Phạm vi TC mới:"],
        ["- Sửa hồ sơ từ danh sách (sheet đầu) — gồm phần Tài sản đảm bảo"],
        ["- Xóa hồ sơ từ danh sách (sheet đầu)"],
        ["- Thanh process trên Chi tiết hồ sơ — highlight theo TRẠNG THÁI HỒ SƠ"],
        [""],
        ["Process steps (UI): Chờ tiếp nhận → Chờ phân công → Chờ thẩm định → Chờ ký duyệt → Đã ký duyệt"],
        ["Assumption/TBD: message success/validate nguyên văn; rule Sửa/Xóa theo trạng thái — confirm SRS."],
    ]
    for r in notes:
        ws0.append(r)
    ws0["A1"].font = Font(bold=True, size=14, color="C0392B")
    ws0.column_dimensions["A"].width = 110
    for row in ws0.iter_rows(min_row=1, max_row=len(notes), min_col=1, max_col=1):
        for c in row:
            c.fill = NOTE_FILL
            c.alignment = WRAP

    # --- Inventory ---
    ws1 = wb.create_sheet("Field_Inventory")
    ws1.append(["#", "Khu vực", "Field / Control", "Type", "Required", "Nguồn UI", "Pattern TC"])
    inv = [
        [1, "List sheet đầu", "Action Sửa", "button/icon", "—", "UI list", "Edit block"],
        [2, "List sheet đầu", "Action Xóa", "button/icon", "—", "UI list", "Delete block"],
        [3, "Form Sửa", "Hình thức (Trực tiếp/Trực tuyến)", "radio", "Y", "UI edit", "Edit"],
        [4, "Form Sửa", "Loại giấy tờ người nộp", "dropdown", "Y", "UI edit", "Edit/required"],
        [5, "Form Sửa", "Số định danh người nộp", "text", "Y", "UI edit", "Edit/required"],
        [6, "Form Sửa", "Họ tên người nộp", "text", "Y", "UI edit", "Edit/required"],
        [7, "Form Sửa / Tài sản", "Tài sản đảm bảo", "textarea", "Y", "UI §5", "Edit asset"],
        [8, "Form Sửa / Tài sản", "Mô tả chi tiết tài sản đảm bảo", "textarea", "Y", "UI §5", "Edit asset"],
        [9, "Form Sửa / Tài sản", "Hiện trạng tài sản", "textarea", "Y", "UI §5", "Edit asset"],
        [10, "Form Sửa / Tài sản", "Tính pháp lý tài sản", "textarea", "Y", "UI §5", "Edit asset"],
        [11, "Chi tiết", "Thanh process — 5 trạng thái", "stepper", "—", "UI detail", "Process bar"],
        [12, "Chi tiết", "Tab CHI TIẾT / THẨM ĐỊNH / QUÁ TRÌNH XỬ LÝ", "tabs", "—", "UI detail", "Process sticky"],
    ]
    for row in inv:
        ws1.append(row)
    for cell in ws1[1]:
        style_cell(cell, HEADER_FILL, HEADER_FONT)
    for col in range(1, 8):
        ws1.column_dimensions[get_column_letter(col)].width = 28

    # --- Plan ---
    ws2 = wb.create_sheet("Plan_TC")
    ws2.append(["Section", "Est. TC", "Patterns", "Ghi chú"])
    plan = [
        ["Sửa hồ sơ (sheet đầu)", 10, "Edit UI + required + asset + status rule", "NEW yellow"],
        ["Xóa hồ sơ (sheet đầu)", 5, "Confirm / success / status / cascade asset", "NEW yellow"],
        ["Chi tiết — Process bar", 11, "5 steps + per-status highlight + sync list", "NEW yellow"],
        ["TỔNG", 26, "", "User đã yêu cầu update — gen trực tiếp"],
    ]
    for row in plan:
        ws2.append(row)
    for cell in ws2[1]:
        style_cell(cell, HEADER_FILL, HEADER_FONT)
    for col in range(1, 5):
        ws2.column_dimensions[get_column_letter(col)].width = 40

    # --- Main TC sheet ---
    ws = wb.create_sheet("TC_MOI", 1)
    ws.append(COLS)
    for cell in ws[1]:
        style_cell(cell, HEADER_FILL, HEADER_FONT)

    for item in rows():
        ws.append([item[c] for c in COLS])
        r = ws.max_row
        is_section = item["Flag"] == "SECTION"
        fill = SECTION_FILL if is_section else YELLOW
        font = SECTION_FONT if is_section else None
        for col in range(1, len(COLS) + 1):
            style_cell(ws.cell(r, col), fill, font)

    widths = [18, 14, 45, 40, 55, 55, 18, 35, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"

    # legend
    ws3 = wb.create_sheet("Legend")
    ws3.append(["Màu", "Ý nghĩa"])
    ws3.append(["Vàng (FFFF99)", "TC MỚI — vừa thêm; cần execute; Kết quả test để trống"])
    ws3.append(["Xanh nhạt", "Dòng SECTION header"])
    ws3.append(["Cam nhạt (sheet Hướng dẫn)", "Hướng dẫn merge — không phải TC"])
    ws3["A2"].fill = YELLOW
    ws3["A3"].fill = SECTION_FILL
    ws3["A4"].fill = NOTE_FILL
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 70

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Wrote {path}")


if __name__ == "__main__":
    targets = [
        Path("/workspace/output/Quan_ly_cho_vay_fix_comment_TC_MOI.xlsx"),
        Path("/opt/cursor/artifacts/HATC_Output/Quan_ly_cho_vay_fix_comment_TC_MOI.xlsx"),
    ]
    for t in targets:
        build(t)
